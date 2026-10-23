"""Borradores personales para validar y guía de defensa basada en evidencia local."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'documentos/Reflexiones_Borrador_y_Guia_Sustentacion.docx'
reflections = {
    'Juan David': (
        'Mi parte fue continuar los cambios del código y las pruebas que faltaban, con ayuda de Codex. '
        'Un cambio que puedo explicar es el de las fotos: para una tarjeta pequeña no hace falta cargar una imagen tan grande como la del detalle. '
        'En la prueba, las cuatro miniaturas pasaron de ocupar 4 MiB a 1 MiB. '
        'También se ajustó el fondo de los botones del sistema para que se vieran mejor en Android 7. '
        'Para mí, lo más importante de este trabajo es entender que compilar no alcanza para decir que todo funciona. '
        'Hay que entrar, guardar los datos y volver a consultarlos. '
        'Las pruebas se hicieron con Firebase local, así que todavía debo tener clara esa diferencia cuando explique los resultados del proyecto.'
    ),
    'Daniel Segura': (
        'Mi aporte fue la revisión de las versiones y las herramientas del proyecto. '
        'En esa parte hay varios nombres que se parecen, pero cumplen funciones distintas, como minSdk, targetSdk y compileSdk. '
        'Por eso es importante explicar qué significa cada uno y mirar los valores que tiene realmente la aplicación. '
        'También está la diferencia entre el JDK que usa Gradle y la versión de Java con la que se compila el código. '
        'Algo que me parece importante es que una biblioteca diga que admite una versión antigua no significa que ya tengamos todas sus funciones probadas. '
        'Hay que comprobarlo al usar la app. '
        'Para la exposición debo repasar la matriz y relacionar esos datos con las pruebas que quedaron guardadas en Android 7 y Android 15.'
    ),
    'Juan Pablo': (
        'Mi parte fue investigar los posibles problemas y comparar formas de resolverlos. '
        'Un ejemplo fue el tema de las fechas: el error de java.time no correspondía a nuestra app, porque el código usa Calendar y DateFormat. '
        'Por eso se corrigió esa parte del informe. '
        'En cambio, el problema de contraste de los botones del sistema sí se veía en Android 7. '
        'La solución fue usar un fondo más oscuro en esa versión y mantener el aspecto del Android nuevo. '
        'También se comparó mejorar la carga de las fotos con agregar otra biblioteca. '
        'De este trabajo me queda que hay que revisar las propuestas de la IA con el código y las pruebas. '
        'Para sustentarlo debo poder explicar por qué se escogió una solución y qué cosas todavía faltan comprobar.'
    ),
    'Ricardo': (
        'La parte que tenía asignada era comprobar los recorridos de la aplicación en Android 7 y Android 15 y organizar las evidencias. '
        'El objetivo de esa tarea es revisar que se pueda iniciar sesión, registrar una mascota y crear un reporte, además de consultar lo que se guardó. '
        'En los resultados del grupo también quedó la prueba de un reporte sin conexión. '
        'El registro conservó su identificador al cerrar la app y se confirmó al volver la conexión, sin duplicarse. '
        'Para explicar esta parte debo revisar las capturas y los resultados de cada versión. '
        'Lo importante es distinguir una pantalla abierta de una función comprobada. '
        'También debo aclarar que los recorridos se hicieron con Firebase local y que eso no demuestra cómo responde el servidor remoto.'
    ),
}
for name, text in reflections.items():
    count = len(text.split())
    assert 100 <= count <= 150, (name, count)
    print(f'{name}: {count} palabras')

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
sec.top_margin = sec.bottom_margin = Cm(1.8)
sec.left_margin = sec.right_margin = Cm(2)
for name, size in [('Normal', 11), ('Title', 20), ('Subtitle', 11), ('Heading 1', 15), ('Heading 2', 12)]:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 1.07
    style.paragraph_format.space_after = Pt(7)
    for border in list(style.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
for name in ('Heading 1', 'Heading 2'):
    doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles[name].paragraph_format.space_before = Pt(10)

def p(text='', style=None):
    return doc.add_paragraph(text, style)
def h(text, level=1):
    return p(text, f'Heading {level}')
def new_page(title):
    doc.add_page_break()
    h(title)
def reflection(name):
    h(name, 2)
    p(reflections[name])

p('Reflexiones y preparación de la sustentación', 'Title')
p('Mascotas Municipales  ·  Taller de compatibilidad Android', 'Subtitle')
p('Estos textos son una base para que cada uno escriba su reflexión. Antes de pasarlos al informe, hay que leerlos y cambiar lo que no corresponda a lo que hizo o aprendió cada persona. Deben quedar entre 100 y 150 palabras. Después está la guía para practicar la exposición con los resultados que ya tenemos.')
h('Borradores de reflexión individual')
reflection('Juan David')
reflection('Daniel Segura')

new_page('Borradores de reflexión individual del grupo')
reflection('Juan Pablo')
reflection('Ricardo')
h('Revisión antes de incorporar al informe', 2)
p('Cada uno debe revisar su párrafo y agregar algo propio que pueda explicar. Si cambia el texto, hay que volver a contar las palabras. La parte de Ricardo describe la tarea que tenía asignada y los resultados del grupo; no afirma que él haya ejecutado las pruebas.')
p('Cuando todos hayan revisado su parte, pueden copiar las cuatro reflexiones al final del informe. La guía puede quedar aparte para estudiar. Lo que toca es entender el trabajo y poder explicar las pruebas con palabras propias.')

new_page('Explicación del proyecto para la sustentación')
h('Presentación de aproximadamente un minuto', 2)
p('Trabajamos sobre Mascotas Municipales para comprobar tres funciones en Android 7 y Android 15: iniciar sesión, registrar y consultar una mascota, y crear y consultar un reporte. Conservamos las versiones de SDK y las bibliotecas del proyecto. Optimizamos la resolución de las miniaturas sin reducir la del detalle, y corregimos el contraste de la barra del sistema en Android 7. Los tres recorridos pasaron con el mismo APK de pruebas locales y datos ficticios. También se comprobó que un reporte pendiente conserva su ID al cerrar el proceso y se confirma sin duplicarse al reconectar. El APK normal se instaló, abrió y pasó la prueba de fotos en ambas versiones. Los recorridos completos se hicieron contra Firebase local, por eso no los presentamos como una prueba del servidor remoto.')
h('Un cambio concreto para explicar', 2)
p('La barra del sistema antes de API 26 solo permite los iconos claros usados en este caso. Con el fondo claro anterior perdían contraste. El código consulta la versión y usa el turquesa oscuro de la paleta cuando esos iconos oscuros no están disponibles. Desde API 26 conserva el fondo claro y solicita iconos oscuros.')
code = '''val darkNavigationIcons = Build.VERSION.SDK_INT >= Build.VERSION_CODES.O
window.navigationBarColor = if (darkNavigationIcons) bg else tealDark
WindowInsetsControllerCompat(window, window.decorView).apply {
    isAppearanceLightStatusBars = false
    isAppearanceLightNavigationBars = darkNavigationIcons
}'''
para = p(code)
para.paragraph_format.space_before = Pt(3)
for run in para.runs:
    run.font.name = 'Consolas'
    run.font.size = Pt(8.5)
p('Archivo: app/src/main/java/com/mascotasmunicipales/MainActivity.kt, función applySystemBars. La mejora de imágenes está en setPhoto: calcula el muestreo para la tarjeta y distingue foto y muestreo en la caché. El detalle conserva la resolución anterior.')
h('Evidencia que se puede mostrar', 2)
p('En tests/evidencias-compatibilidad/funcionales-2026-10-22/taller22d, abrir api24 y api35. entorno.json identifica la versión y el SHA-256; run.json y recover.json registran los resultados. Las capturas 03 y 05 muestran mascota y reporte; 10 y 11 muestran recuperación y confirmación. servidor-final.json registra una mascota y dos reportes por cuenta, uno normal y otro recuperado.')

new_page('Preguntas y preparación de la defensa')
questions = [
    ('Por qué no bajaron minSdk ni las bibliotecas', 'La app ya declara minSdk 24, que incluye Android 7. Los recorridos elegidos funcionaron en API 24 y 35; no hubo evidencia que justificara retroceder las versiones.'),
    ('Qué significan los tres SDK', 'minSdk es el mínimo instalable, targetSdk declara el nivel de comportamiento objetivo y compileSdk permite compilar contra sus API. En el proyecto se mantienen 24, 33 y 35, respectivamente.'),
    ('Qué demuestra la mejora de fotos', 'En la prueba aislada los cuatro bitmaps de miniaturas pasaron de 4 MiB a 1 MiB. Es una reducción del 75 % de esos bitmaps; no de toda la memoria ni del tiempo de ejecución de la app.'),
    ('Qué consumió más memoria en el recorrido', 'El mapa tuvo mayor mediana PSS que el directorio: 67,7 frente a 63,9 MiB en Android 7 y 121,4 frente a 116,2 MiB en Android 15. Son observaciones con instrumentación y caché, no un benchmark en teléfonos físicos.'),
    ('Cómo saben que no hubo duplicación', 'El reporte pendiente no existía en el servidor sin red; conservó su ID tras detener el proceso y se confirmó después de reconectar. La cuenta terminó con dos reportes: el normal y el recuperado. El botón quedó bloqueado durante el envío.'),
    ('Qué se probó realmente y qué quedó fuera', 'El mismo APK QA pasó los tres recorridos contra Authentication y Firestore locales en ambos Android. El APK normal es distinto. No se verificaron estos recorridos contra producción ni todos los caminos de QR, accesibilidad y roles profesionales.'),
]
for question, answer in questions:
    para = p()
    para.add_run(question + '. ').bold = True
    para.add_run(answer)
h('Si el profesor cambia una condición', 2)
p('Conservar una biblioteca: los ajustes de muestreo y contraste ya mantienen las bibliotecas. Evitar una biblioteca archivada: osmdroid requeriría evaluar una sustitución, conservar mapa y atribución y repetir pruebas; esa migración no está hecha. Conservar una función: demostrar su recorrido completo antes de aceptar una alternativa que pueda afectarla.')
h('Antes de entregar y exponer', 2)
p('Revisar las cuatro reflexiones y agregarlas al Word; guardar el informe actualizado en el paquete de entrega. Tener a mano el proyecto, ambos entorno.json, capturas y verificaciones del servidor. Explicar el uso de Codex y distinguir pruebas automatizadas de pruebas manuales. El trabajo Offline First con GUARDAR local y SINCRONIZAR manual hacia RTDB sigue pendiente y separado de este taller.')

doc.core_properties.title = 'Reflexiones y preparación de la sustentación'
doc.core_properties.author = 'Grupo Mascotas Municipales'
doc.core_properties.subject = 'Borradores individuales para revisión y guía basada en pruebas locales'
doc.save(OUT)
print(OUT)
