from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(r'D:\MascotasMunicipalesNative_Corregido')
OUT = ROOT / 'documentos' / 'Taller_Compatibilidad_Android_Avance_Grupo.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.7)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.header_distance = sec.footer_distance = Cm(.8)
for name, size in [('Normal', 10.5), ('Title', 21), ('Subtitle', 11.5), ('Heading 1', 14), ('Heading 2', 11.5)]:
    s = doc.styles[name]
    s.font.name = 'Calibri'
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.line_spacing = 1.07
for name in ['Heading 1', 'Heading 2']:
    doc.styles[name].paragraph_format.space_before = Pt(9)
    doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles[name].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(7)
for border_node in list(doc.styles['Title'].element.xpath('.//w:pBdr')):
    border_node.getparent().remove(border_node)
header = sec.header.paragraphs[0]
header.text = 'Mascotas Municipales  |  Taller de compatibilidad Android'
header.runs[0].font.size = Pt(8)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Avance del 20 de octubre de 2026  ·  Página ').font.size = Pt(8)
for instruction in ['PAGE', 'NUMPAGES']:
    if instruction == 'NUMPAGES':
        footer.add_run(' de ').font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), instruction)
    footer._p.append(field)

def p(text='', bold=False):
    para = doc.add_paragraph()
    para.add_run(text).bold = bold
    return para

def h(text, level=1):
    return doc.add_paragraph(text, f'Heading {level}')

def page(title):
    doc.add_page_break()
    h(title)

def bullet(text):
    para = doc.add_paragraph(text, 'List Bullet')
    para.paragraph_format.space_after = Pt(4)
    return para

def table(headers, rows, widths, font=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Cm(width)
    all_rows = [headers] + list(rows)
    for i, values in enumerate(all_rows):
        row = t.rows[0] if i == 0 else t.add_row()
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        trpr.append(no_split)
        if i == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for cell, text, width in zip(row.cells, values, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(text)
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'bottom', 'left', 'right']:
                node = OxmlElement('w:' + edge)
                node.set(qn('w:w'), '75' if edge in ['top', 'bottom'] else '100')
                node.set(qn('w:type'), 'dxa')
                margins.append(node)
            tcpr.append(margins)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge)
                for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    e.set(qn('w:' + key), value)
                borders.append(e)
            tcpr.append(borders)
            fill = 'E7EEF5' if i == 0 else ('F7F9FB' if i % 2 == 0 else 'FFFFFF')
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), fill)
            tcpr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.02
                for run in para.runs:
                    run.font.size = Pt(font)
                    run.bold = i == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def link(label, url):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(3)
    rel = para.part.relate_to(url, RT.HYPERLINK, is_external=True)
    a = OxmlElement('w:hyperlink')
    a.set(qn('r:id'), rel)
    r = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '194E78'); props.append(color)
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '18'); props.append(size)
    r.append(props)
    text = OxmlElement('w:t'); text.text = label; r.append(text)
    a.append(r); para._p.append(a)

doc.add_paragraph('Compatibilidad Android de Mascotas Municipales', 'Title')
doc.add_paragraph('Informe de avance del taller', 'Subtitle')
p('Integrantes: Juan David, Juan Pablo, Daniel Segura y Ricardo.')
p('Asignatura: Desarrollo Móvil  ·  Fecha: 20 de octubre de 2026')
h('1 Lo que llevamos hasta ahora')
p('En este taller estamos trabajando con nuestra aplicación Mascotas Municipales. El objetivo es conservar sus funciones y comprobar que el mismo APK funcione en Android 7.0 y Android 15. No estamos creando una aplicación nueva ni cambiando versiones solo para que la compilación pase.')
p('Hasta este momento revisamos la configuración del proyecto, corregimos la información de los aportes individuales y aplicamos una mejora en la carga de las fotos. El APK inicial y el modificado se instalaron y abrieron en los dos emuladores. La prueba de miniaturas pasó en ambos. Todavía falta completar los recorridos de las tres funciones elegidas, por eso presentamos este documento como avance.')
h('Distribución del trabajo', 2)
table(['Integrante', 'Responsabilidad acordada'], [
    ['Juan David', 'Cambios de código, rama de trabajo, registro técnico y comprobación de la mejora.'],
    ['Daniel Segura', 'Auditoría de SDK, herramientas y dependencias; consolidación de la matriz técnica.'],
    ['Juan Pablo', 'Investigación de riesgos, comparación de alternativas y registro del uso de IA.'],
    ['Ricardo', 'Pruebas completas de las funciones en ambos Android y organización de sus evidencias.'],
], [3.1, 14.3], 9.5)
h('Las tres funciones que vamos a conservar', 2)
bullet('Inicio de sesión: entrar con una cuenta ciudadana ficticia y comprobar que se prepara el perfil.')
bullet('Registro y consulta de mascotas: guardar una ficha, confirmar la respuesta del servidor y volver a consultarla.')
bullet('Creación y consulta de reportes: enviar un reporte ficticio, confirmar su guardado y encontrarlo en la cuenta correspondiente.')
p('El mapa, QR, accesibilidad y los flujos profesionales se conservan. No los eliminamos por no estar dentro de estos tres recorridos principales. Firebase sigue sin cambios de reglas ni servicios nuevos; no se habilitó facturación.')

page('2 Auditoría y matriz técnica')
p('La revisión se hizo sobre los archivos reales del proyecto, el APK generado y los manifiestos de las bibliotecas que resolvió Gradle. Estas son las versiones encontradas, no versiones propuestas para actualizar.')
table(['Elemento', 'Valor comprobado', 'Decisión y evidencia'], [
    ['minSdk', '24 · Android 7.0', 'Se conserva. app/build.gradle.kts y APK.'],
    ['targetSdk', '33 · Android 13', 'Se conserva para este cambio. No convierte Android 15 en Android 13.'],
    ['compileSdk', '35 · Android 15', 'Se conserva como SDK para compilar. No aumenta por sí solo el mínimo instalable.'],
    ['AGP / Gradle', '9.3.1 / 9.5.0', 'Configuración raíz y wrapper. La compilación actual pasó.'],
    ['JDK del lanzador', '21.0.6', 'Salida de gradlew --version del entorno usado.'],
    ['JDK del daemon', '25', 'toolchainVersion=25 y registro del daemon con javaVersion=25.'],
    ['Nivel Java del código', '11', 'compileOptions. No es el JDK que ejecuta Gradle ni garantiza todas las API en Android 7.'],
    ['Google Services', '4.5.0', 'Plugin que procesa la configuración de cliente de Firebase.'],
    ['Versión de la app', '1.5 · versionCode 6', 'Se mantiene en el APK de esta comparación.'],
], [3.3, 3.4, 10.7], 9)
h('Dependencias y mínimos encontrados', 2)
table(['Biblioteca o componente', 'Versión', 'API mínima del manifest local'], [
    ['Firebase Authentication', '24.2.0', '23'],
    ['Cloud Firestore', '26.6.0', '23'],
    ['Google Code Scanner', '16.1.0', '21 declarada; la guía actual pide 23 o superior.'],
    ['osmdroid', '6.1.20', '1 declarada; no la tomamos como garantía de soporte real de Android 1.'],
    ['AndroidX Core KTX / AppCompat', '1.15.0 / 1.7.1', '21 / 21'],
    ['Lifecycle ViewModel / LiveData Core', '2.6.2 / 2.6.2', '14 / 14'],
], [6.6, 3.2, 7.6], 9)
p('Firebase BoM 34.19.0 coordina versiones; no es una biblioteca ejecutable con minSdk propio. La fusión del manifest pasó y el APK declara API 24. Eso permite instalarlo, pero no demuestra todas las funciones de cada biblioteca. Authentication y Firestore no exigen obligatoriamente Google Play Services en estas versiones; el lector QR sí tiene esa dependencia de ejecución. Fuentes [1], [4] y [5].')

page('3 Riesgos y alternativas investigadas')
h('Diferencia real de interfaz entre versiones', 2)
p('En las capturas vimos que Android 7 muestra los botones del sistema claros sobre un fondo también claro. El código fija ese fondo y solicita iconos oscuros con isAppearanceLightNavigationBars. Según Android, esa solicitud no tiene efecto antes de API 26. Por eso registramos un problema de contraste para revisar. La aplicación abre; no presentamos este hallazgo como un cierre de la app. Fuente [2].')
table(['Criterio', 'A Fondo según versión', 'B Fondo oscuro en ambas versiones'], [
    ['Propuesta', 'Antes de API 26 usar fondo oscuro; desde API 26 conservar fondo claro con iconos oscuros.', 'Usar barra oscura e iconos claros en los dos Android.'],
    ['Ventaja', 'Conserva el aspecto actual donde los iconos oscuros están disponibles.', 'Reduce las diferencias de configuración entre versiones.'],
    ['Riesgo y mantenimiento', 'Hay que probar las dos ramas y los modos de accesibilidad.', 'Cambia también el aspecto del Android moderno y requiere revisión visual.'],
    ['Estado', 'Propuesta preferida para discutir y probar. No está implementada.', 'Alternativa comparada. No está implementada.'],
], [3, 7.2, 7.2], 9.5)
p('Las dos alternativas mantienen la misma aplicación y sus funciones. Ninguna exige retroceder bibliotecas. La preferencia por A todavía debe revisarse con el equipo; la mejora ya aplicada en este avance es la de memoria descrita en la siguiente sección.')
h('Otros riesgos que sí corresponden al proyecto', 2)
bullet('QR: depende de Google Play Services y de su módulo. Ya existe la consulta manual como respaldo; falta probar ambos caminos en los dos Android. Fuente [4].')
bullet('Mapa: osmdroid está archivado y las teselas necesitan red. Un fallo de certificados en Android 7 es una hipótesis que debe comprobarse con Logcat; no lo hemos observado. Fuente [7].')
bullet('Cambio de targetSdk: subirlo a 35 activa nuevas exigencias de interfaz en Android 15, entre ellas el comportamiento de borde a borde. No lo cambiamos para resolver un riesgo que no requería esa actualización. Fuente [6].')
bullet('Firebase remoto: README señala despliegues de reglas e índices pendientes. No consultamos la consola remota en esta revisión ni asumimos que esté actualizada.')
h('Corrección del ejemplo sobre fechas', 2)
p('El ejemplo de java.time del aporte inicial no describe nuestra app. El código usa Calendar y DateFormat; no encontramos LocalDate ni LocalDateTime. Tampoco activamos coreLibraryDesugaring ni tenemos un Logcat con ese NoClassDefFoundError. Por eso quitamos esa afirmación del informe.')

page('4 Cambio aplicado y resultados')
p('El cambio se preparó en la rama codex/android-compatibilidad, creada desde el commit 4764e19. Antes de comenzar no había modificaciones en archivos versionados. Se conservaron los documentos y carpetas locales del equipo, y se guardó un APK inicial para comparar.')
h('Qué se cambió y por qué', 2)
p('Las cuatro fotos de demostración tienen 1024 × 1024 píxeles. Antes se decodificaban siempre a 512 × 512, incluso en tarjetas de 76 dp. Ajustamos setPhoto en MainActivity para que las miniaturas carguen una resolución acorde con la tarjeta. La caché identifica tanto la foto como su muestreo, de manera que una miniatura no sustituya la imagen del detalle. El detalle conserva la resolución anterior. Fuente [3].')
p('Comparamos ajustar el muestreo existente con agregar una biblioteca de imágenes. Aplicamos el primer camino porque es un cambio pequeño para cuatro recursos locales y no agrega dependencias. Agregar otra biblioteca podría ayudar con carga asíncrona, pero ampliaría el mantenimiento y las pruebas necesarias.')
table(['Prueba', 'Android 7.0 · API 24', 'Android 15 · API 35'], [
    ['Instalación y arranque del APK inicial', 'Correctos', 'Correctos'],
    ['Instalación y arranque del APK modificado', 'Correctos', 'Correctos'],
    ['Cuatro fotos con muestreo anterior', '4.194.304 bytes · 4 MiB', '4.194.304 bytes · 4 MiB'],
    ['Cuatro miniaturas con el cambio', '1.048.576 bytes · 1 MiB', '1.048.576 bytes · 1 MiB'],
    ['Reducción en esos bitmaps', '75 %', '75 %'],
    ['Caché y resolución del detalle', 'Prueba aprobada', 'Prueba aprobada'],
    ['Inicio de sesión completo', 'Pendiente', 'Pendiente'],
    ['Registro y consulta de mascota', 'Pendiente', 'Pendiente'],
    ['Creación y consulta de reporte', 'Pendiente', 'Pendiente'],
], [7.2, 5.1, 5.1], 9.5)
p('La instrumentación ejecutó el decodificador real de cada Android. En ambos emuladores, 76 dp equivalen a 152 píxeles. La comparación reproduce el muestreo anterior y comprueba la caché y las dimensiones del detalle. No mide un recorrido completo ni demuestra que toda la app use 75 % menos memoria.')
h('Compilación e incidencias', 2)
p('Pasaron assembleDebug, lintDebug y assembleDebugAndroidTest. El último informe de lint registra 0 errores y 40 advertencias; no lo presentamos como una aplicación libre de advertencias. No apareció un error de compilación de la app. Hubo problemas de descarga, registro del SDK y ruta de los AVD; se resolvieron al preparar el entorno y no se atribuyen a incompatibilidades del código.')

page('5 Evidencia en los dos Android')
p('Estas capturas corresponden al mismo APK modificado. Muestran que la aplicación abrió el formulario de acceso. No prueban que ya se haya iniciado sesión ni que se hayan guardado mascotas o reportes.')
t = table(['Android 7.0 · API 24', 'Android 15 · API 35'], [['', '']], [8.7, 8.7], 10)
for cell, api in zip(t.rows[1].cells, [24, 35]):
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = para.add_run().add_picture(str(ROOT / f'tests/evidencias-compatibilidad/api{api}/inicio.png'), width=Cm(6.7))
    picture._inline.docPr.set('descr', f'Formulario de acceso real de Mascotas Municipales en Android API {api}, sin datos personales.')
p('Entorno: emuladores x86_64 de 720 × 1280 píxeles, densidad 320 y dos núcleos. Android 7 se configuró con 1536 MiB de RAM y Android 15 con 2048 MiB. No son mediciones en teléfonos físicos de gama baja.')
p('Archivo utilizado: apk/MascotasMunicipales-compatibilidad.apk', True)
p('SHA-256 del APK modificado:')
para = p('5C7462623551FD6C960C05C0FC4BB3DB8912CA7AB5E21EE18A60537712BD19C8')
for run in para.runs:
    run.font.name = 'Consolas'; run.font.size = Pt(9)
p('Los archivos entorno.json de API 24 y API 35 registran exactamente esta misma huella.')
p('Evidencia del proyecto: tests/evidencias-compatibilidad/api24 y api35 contienen instalacion.txt, inicio.txt, inicio.png, memoria-inicio.txt, fotos.txt y entorno.json. Las carpetas api24-inicial y api35-inicial conservan la comparación de instalación y arranque previa al cambio.')

page('6 Respuestas técnicas de la primera parte')
h('Pregunta 1 Mínimos de la app y sus dependencias', 2)
p('La app declara minSdk 24 y el APK confirma ese valor. En los manifiestos locales de las bibliotecas encontramos API 23 para Authentication y Firestore; API 21 para Code Scanner, AppCompat y Core KTX; y API 14 para las dos bibliotecas Lifecycle. En el AAR de osmdroid aparece 1, pero ese dato aislado no permite asegurar funcionamiento real en versiones tan antiguas. La guía de Code Scanner, además, pide API 23 o superior. Dejamos esa diferencia registrada en la matriz.')
p('Lo comprobamos con app/build.gradle.kts, el informe de fusión del manifest, los manifiestos de los AAR resueltos y la inspección del APK. La fusión y la compilación pasaron, y el APK se instaló en API 24. Todavía falta probar todos los caminos de las bibliotecas dentro de los recorridos elegidos; no basta con mirar el mínimo declarado. Fuentes [1], [4] y [5].')
h('Pregunta 2 Versiones que conservamos y motivo', 2)
p('Conservamos compileSdk 35 porque permite compilar contra esas API, sin obligar a instalarlas en el Android antiguo. Conservamos targetSdk 33 para no mezclar la mejora de memoria con cambios de comportamiento de la plataforma. No significa que Android 15 se comporte por completo como Android 13. Una actualización del objetivo a 35 necesitaría revisar también los márgenes y las barras del sistema. Fuentes [1] y [6].')
p('También conservamos AGP 9.3.1, Gradle 9.5.0 y las bibliotecas actuales porque la compilación pasó y no encontramos una razón comprobada para retrocederlas. Diferenciamos el JDK del lanzador, 21.0.6, del daemon que ejecutó la compilación, 25. El nivel Java 11 del código es otra configuración distinta; por sí solo no garantiza que cualquier API de Java esté disponible en Android 7.')
h('Pregunta 3 Diferencias de API e implementación compatible', 2)
p('El ejemplo que sí encontramos está en la barra de navegación del sistema: isAppearanceLightNavigationBars permite pedir iconos oscuros, pero no funciona antes de API 26. En Android 7 observamos iconos claros sobre el fondo claro de la app. Para mantener una sola aplicación proponemos condicionar el fondo por versión: oscuro en API 24 y conservar el comportamiento actual donde los iconos oscuros sí están disponibles. Comparamos esta idea con usar fondo oscuro en ambas versiones. La corrección todavía está pendiente; no la presentamos como implementada. Fuente [2].')
p('La app ya separa recursos de tema en values y values-v31, que Android selecciona según la versión. También vimos una diferencia en el dibujo de la huella del encabezado, relacionada con la fuente de emoji del sistema. Esa variación visual no es evidencia de una función eliminada o de un cierre.')

page('7 Respuestas técnicas y reproducción')
h('Pregunta 4 Recursos y optimización concreta', 2)
p('Todavía no podemos afirmar qué operación consume más recursos en toda la aplicación. Falta comparar el mapa y los directorios durante su uso real. Lo que sí identificamos fue un gasto evitable en la carga de las fotos de las tarjetas. En la prueba aislada de ambos Android, los cuatro bitmaps pasaron de 4 MiB a 1 MiB, sin reducir las dimensiones del detalle. Esa es la optimización comprobada. No medimos todavía CPU, batería ni fluidez completa. Fuente [3].')
h('Pregunta 5 Riesgos de retroceder bibliotecas', 2)
p('Retroceder una biblioteca puede dejar fuera correcciones de seguridad, cambios de mantenimiento o funciones que usa el proyecto; también puede crear conflictos con otras dependencias. El efecto depende de la versión concreta, así que no afirmamos que cualquier versión anterior sea insegura. En este avance no retrocedimos Firebase, osmdroid ni Code Scanner.')
p('Para las imágenes descartamos añadir una biblioteca nueva y preferimos ajustar el muestreo existente. Para la interfaz comparamos un fondo condicionado por API con un fondo oscuro común. Tampoco elegimos bajar targetSdk como remedio general, porque eso no agrega una API que falte en Android 7. Subir minSdk dejaría fuera el dispositivo antiguo requerido. osmdroid requiere seguimiento aparte por estar archivado; eso no demuestra una vulnerabilidad concreta en la app. Fuentes [1] y [7].')
h('Pregunta 6 Evidencia del mismo APK y las tres funciones', 2)
p('Los registros de API 24 y API 35 contienen la misma huella SHA-256 del APK modificado. También incluyen instalación, arranque, capturas y la prueba de fotos. Eso demuestra que se usó el mismo archivo para esas comprobaciones. No demuestra todavía que se completaron las tres funciones: faltan los recorridos de inicio de sesión, mascota y reporte en ambos Android. Ricardo debe añadir las evidencias de cada uno antes de dar esta respuesta por terminada.')
h('Cómo repetir lo comprobado', 2)
p('Desde la carpeta del proyecto, con el SDK y los dos emuladores preparados:')
for command in [
    'gradlew.bat :app:assembleDebug :app:lintDebug :app:assembleDebugAndroidTest --offline',
    '.\\scripts\\verify-android-compatibility.ps1 -Serial emulator-5554 -ExpectedApi 24',
    '.\\scripts\\verify-android-compatibility.ps1 -Serial emulator-5556 -ExpectedApi 35',
]:
    para = p(command)
    para.paragraph_format.space_after = Pt(7)
    for run in para.runs:
        run.font.name = 'Consolas'; run.font.size = Pt(9)
p('Los seriales son ejemplos: hay que usar los que muestre adb devices -l. No se debe recompilar entre las dos pruebas del mismo APK. La instrumentación se instala aparte y no forma parte del APK normal.')

page('8 Uso de IA y trabajo pendiente')
p('Usamos Codex como apoyo para revisar el código, contrastar la documentación, preparar la mejora, ejecutar comprobaciones y consolidar este informe. La redacción reúne el avance del grupo; no reemplaza la revisión de cada integrante ni su explicación en la sustentación.')
table(['Consulta o apoyo', 'Qué se aceptó o corrigió'], [
    ['Compatibilidad de fechas', 'Se descartó presentar java.time como un fallo real: el proyecto usa Calendar y DateFormat.'],
    ['SDK y dependencias', 'Se contrastaron versiones y manifiestos. Se corrigió la explicación de targetSdk y la dependencia de Play Services.'],
    ['Memoria de fotografías', 'Se preparó el muestreo por tamaño y una prueba con el decodificador de Android. El resultado quedó en archivos verificables.'],
    ['Redacción y pruebas', 'Se integraron los aportes corregidos y se marcaron los pendientes. No se aceptó presentar pruebas sin evidencia ni reflexiones individuales inventadas.'],
], [5.1, 12.3], 9)
h('Lo que nos falta antes de entregar', 2)
bullet('Completar en ambos Android los tres recorridos funcionales, con cuentas ficticias y confirmación del servidor. Preparar Firebase local como indican los protocolos del proyecto.')
bullet('Revisar con el equipo el contraste de la navegación en Android 7 y decidir si se aplica el ajuste propuesto. Revisar visualmente miniaturas, detalle y accesibilidad.')
bullet('Completar la observación de recursos del mapa y directorio. Las hipótesis de QR y certificados siguen pendientes de prueba.')
bullet('Añadir las cuatro reflexiones individuales de 100 a 150 palabras. Cada integrante debe escribir la suya a partir de su trabajo y de las pruebas que haya revisado o realizado.')
p('No se hicieron despliegues, cambios de reglas, creación de roles remotos ni integración del taller Offline First con RTDB. Ese último requerimiento pertenece a otro trabajo y sigue pendiente.')
h('Fuentes consultadas y archivos de apoyo', 2)
link('[1] Android Developers · SDK mínimo y objetivo', 'https://developer.android.com/guide/topics/manifest/uses-sdk-element')
link('[2] Android Developers · WindowInsetsControllerCompat', 'https://developer.android.com/reference/androidx/core/view/WindowInsetsControllerCompat')
link('[3] Android Developers · Carga eficiente de bitmaps', 'https://developer.android.com/topic/performance/graphics/load-bitmap')
link('[4] Google Developers · Google Code Scanner', 'https://developers.google.com/ml-kit/vision/barcode-scanning/code-scanner')
link('[5] Firebase · Dependencias de Google Play Services', 'https://firebase.google.com/docs/android/android-play-services')
link('[6] Android Developers · Cambios de Android 15 para target 35', 'https://developer.android.com/about/versions/15/behavior-changes-15')
link('[7] Repositorio oficial de osmdroid', 'https://github.com/osmdroid/osmdroid')
p('Archivos del proyecto: README.md, ROLES.md, archivos Gradle, MainActivity.kt, estilos, informe de fusión del manifest, android-compatibility.md y evidencias-compatibilidad. También se revisaron el enunciado del taller y los aportes de Juan Pablo y Daniel. Las fuentes en línea se consultaron el 20 de octubre de 2026.')

doc.core_properties.title = 'Compatibilidad Android de Mascotas Municipales'
doc.core_properties.subject = 'Informe de avance del taller con auditoría y pruebas parciales'
doc.core_properties.author = 'Equipo Mascotas Municipales'
doc.core_properties.keywords = 'Android 7, Android 15, Kotlin, compatibilidad, avance'
doc.core_properties.comments = 'Informe consolidado con apoyo de Codex y evidencias disponibles al 20 de octubre de 2026.'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
