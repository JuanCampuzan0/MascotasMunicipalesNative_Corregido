"""Actualiza una copia del informe aportado por el grupo con evidencias verificadas."""
from pathlib import Path
import argparse
import json
from statistics import median
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser()
parser.add_argument('source', type=Path)
args = parser.parse_args()
BASE = ROOT / 'tests/evidencias-compatibilidad/funcionales-2026-10-22'
EVIDENCE = BASE / 'taller22d'
def read(path):
    return json.loads(path.read_text(encoding='utf-8-sig'))
runs = {api: read(EVIDENCE / f'api{api}/run.json') for api in (24, 35)}
envs = {api: read(EVIDENCE / f'api{api}/entorno.json') for api in (24, 35)}
for api in runs:
    assert runs[api]['result'] == 'APROBADO'
    assert read(EVIDENCE / f'api{api}/recover.json')['result'] == 'APROBADO'
    assert read(EVIDENCE / f'api{api}/servidor-final.json')['ownerReportCount'] == 2
assert envs[24]['apkSha256'] == envs[35]['apkSha256']
qa_hash = envs[24]['apkSha256']
normal_hash = read(BASE / 'normal-api24/entorno.json')['apkSha256']
assert normal_hash == read(BASE / 'normal-api35/entorno.json')['apkSha256']
doc = Document(args.source)
assert len(doc.paragraphs) == 99 and len(doc.tables) == 7, 'Revisar la estructura de la nueva fuente antes de actualizar.'
assert 'Todavía falta' in doc.paragraphs[6].text
updates = {
1: 'Informe del taller con recorridos verificados',
3: 'Asignatura: Desarrollo Móvil  ·  Actualización: 22 de octubre de 2026',
4: '1 Estado del trabajo',
6: 'Ya comprobamos los tres recorridos en Android 7.0 y Android 15 con el mismo APK de pruebas locales. El registro, ingreso y cierre de sesión funcionaron; también se guardaron y consultaron mascotas y reportes con confirmación del servidor local. Conservamos la mejora de fotos y corregimos el contraste de la navegación en Android 7. Faltan las reflexiones personales y la sustentación de cada integrante.',
13: 'El mapa, QR, accesibilidad y los flujos profesionales se conservan. Para probar sin crear datos remotos añadimos una variante QA con paquete y datos separados. Usa Authentication y Firestore locales; la app normal conserva su configuración. No cambiamos reglas ni índices, ni habilitamos facturación.',
24: 'En Android 7 observamos botones del sistema claros sobre un fondo claro. La solicitud isAppearanceLightNavigationBars no tiene efecto antes de API 26. Aplicamos un fondo turquesa oscuro antes de esa API y conservamos el fondo claro con iconos oscuros desde API 26. Las capturas nuevas muestran la diferencia corregida. Fuente [2].',
26: 'Aplicamos la alternativa A porque corrige el contraste en Android 7 y conserva el aspecto del Android moderno. El color oscuro se obtiene de la paleta que ya usa la app. La validación visual de todos los modos de accesibilidad queda como comprobación adicional.',
29: 'Mapa: osmdroid está archivado y sus teselas necesitan red. El mapa de Zipaquirá se mostró en ambos emuladores durante estos recorridos. No observamos un fallo de certificados en esas pruebas; esto no garantiza todas las redes ni todos los equipos. Fuente [7].',
35: '4 Cambios aplicados y resultados',
36: 'La rama codex/android-compatibilidad conserva el punto de partida 4764e19 y el avance del equipo registrado en a9250e6. Para completar los recorridos añadimos pruebas en una variante local separada y el ajuste de contraste. No sobrescribimos el documento anterior ni eliminamos funciones.',
43: 'Pasaron assembleDebug, lintDebug, assembleDebugAndroidTest y las tareas de compilación y lint de QA. Lint registra 0 errores y 40 advertencias. También pasaron las 33 pruebas de reglas en Firebase local. La variante QA tuvo un conflicto de etiqueta del manifest y una carrera al inicializar Firebase; se corrigieron en su configuración y prueba. Otro intento esperaba Inicio, pero la app restauraba Perfil correctamente; ajustamos el recorrido.',
46: 'Estas capturas muestran las fichas guardadas y consultadas con el mismo APK QA en los dos Android. La instrumentación recorre los controles reales de la interfaz, comprueba la lectura del servidor local y guarda resultados JSON. El APK normal también se instaló, abrió y pasó la prueba de fotos en ambos.',
49: 'APK de los recorridos locales: MascotasMunicipales-QA-local-2026-10-22.apk',
50: 'SHA-256 del mismo APK QA usado en ambos Android:',
51: qa_hash,
52: 'El APK normal es distinto: MascotasMunicipales-compatibilidad-2026-10-22.apk. Conserva Firebase remoto. Los recorridos completos corresponden a QA local; no afirmamos haberlos probado contra producción.',
53: 'Evidencias finales: tests/evidencias-compatibilidad/funcionales-2026-10-22/taller22d/api24 y api35. Incluyen 12 capturas por Android, run.json, recover.json, entorno.json y verificaciones del servidor. normal-api24 y normal-api35 identifican la prueba del APK normal. Las evidencias del 20 de octubre se conservan como antecedentes.',
55: '6 Respuestas técnicas de la primera parte',
58: 'Lo comprobamos con Gradle, los manifests de los AAR resueltos y la inspección del APK. La instalación y los tres recorridos pasaron en API 24 y 35. Esa comprobación no cubre todos los caminos de cada biblioteca: el lector QR con cámara sigue fuera de estos tres recorridos. Fuentes [1], [4] y [5].',
63: 'La barra de navegación muestra una diferencia concreta: los iconos oscuros no están disponibles antes de API 26. La implementación consulta Build.VERSION.SDK_INT y usa turquesa oscuro con iconos claros en Android 7; desde API 26 conserva el fondo claro y solicita iconos oscuros. El mismo código se compiló en los APK usados en ambos Android. Fuente [2].',
68: 'En esta ejecución, el mapa tuvo más memoria proporcional del proceso que el directorio en ambos emuladores. Tomamos tres muestras de Debug.MemoryInfo.totalPss por pantalla, tras esperar la carga. Las medianas se presentan en la sección 9. Es una observación en emuladores con instrumentación, no una clasificación definitiva de consumo ni una medición de CPU o batería. La mejora concreta de miniaturas conserva el resultado aislado de 4 MiB a 1 MiB. Fuente [3].',
73: 'Los dos entorno.json finales contienen la misma huella SHA-256 del APK QA y la identificación de API 24 y 35. run.json registra sesión, mascota y reporte confirmados por el servidor; recover.json registra recuperación y cierre de sesión. Las capturas muestran los detalles y estados. La prueba usa una variante local del mismo proyecto con su código funcional compartido, paquete separado y Firebase emulado; no demuestra el funcionamiento del backend remoto del APK normal.',
74: 'Cómo repetir los recorridos',
75: 'El protocolo completo está en tests/android-functional-compatibility.md. Primero se inician Authentication y Firestore locales con demo-mascotas-compat, se ejecutan las reglas y después se preparan las cuatro fichas ficticias. No ejecutar las reglas mientras se recorren los formularios, porque limpian la base local.',
76: 'gradlew.bat -PcompatibilityTests :app:assembleDebug :app:lintDebug :app:assembleCompatibility :app:assembleCompatibilityAndroidTest :app:lintCompatibility --offline',
77: '.\\scripts\\run-functional-compatibility.ps1 -Serial emulator-5554 -ExpectedApi 24 -RunId nuevo01',
78: '.\\scripts\\run-functional-compatibility.ps1 -Serial emulator-5556 -ExpectedApi 35 -RunId nuevo01',
79: 'Usar los seriales de adb devices y un RunId nuevo en cada repetición. Conservar el mismo APK entre dispositivos. El APK QA requiere los emuladores Firebase locales y no reemplaza el instalador normal.',
81: '8 Uso de IA y cierre del taller',
82: 'Usamos Codex para revisar el código, contrastar documentación, preparar cambios y automatizar pruebas. Juan David continuó la parte de recorridos que estaba asignada a Ricardo. Los resultados provienen de emuladores y archivos verificables; no los presentamos como pruebas manuales realizadas por Ricardo. Cada integrante debe revisar y explicar el trabajo en la sustentación.',
84: 'Lo que falta antes de entregar',
85: 'Revisar en grupo las capturas y resultados de las tres funciones, ya verificados en ambos Android contra Firebase local. Preparar el proyecto modificado o su enlace al repositorio para la entrega.',
86: 'Añadir las cuatro reflexiones individuales de 100 a 150 palabras y practicar la defensa. Cada integrante debe redactar la suya a partir de lo que revisó o realizó.',
87: 'El lector QR con cámara, todos los modos de accesibilidad y los flujos profesionales no se recorrieron de nuevo en esta comparación. Las pruebas remotas y en teléfonos físicos también quedan fuera de lo demostrado.',
88: 'Entregar el informe, APK y proyecto o enlace según lo solicitado por el profesor. El paquete de evidencias distingue el instalador normal del APK QA que necesita Firebase local.',
98: 'Archivos de apoyo: Gradle, MainActivity.kt, CompatibilityApplication.kt, FunctionalCompatibilityInstrumentation.kt, los dos scripts de recorridos locales y los protocolos de tests. Se conservaron las fuentes y aportes del informe anterior. Las fuentes web fueron consultadas el 20 de octubre de 2026; los nuevos recorridos se ejecutaron el 22 de octubre.'
}
for i, text in updates.items():
    para = doc.paragraphs[i]
    para.clear()
    para.add_run(text)

def cell_text(cell, text):
    cell.text = text
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.size = Pt(9)

cell_text(doc.tables[0].rows[1].cells[1], 'Cambios de código, rama, mejora de fotos y continuación de los recorridos con apoyo de automatización.')
cell_text(doc.tables[0].rows[4].cells[1], 'Recorridos inicialmente asignados. Juan David continuó esta parte y dejó las evidencias para revisión del grupo.')
cell_text(doc.tables[3].rows[4].cells[1], 'Aplicada y comprobada en la paleta normal de ambos Android.')
cell_text(doc.tables[3].rows[4].cells[2], 'Comparada y descartada para conservar el aspecto moderno.')
for row in doc.tables[4].rows[7:]:
    cell_text(row.cells[1], 'Aprobado en QA local')
    cell_text(row.cells[2], 'Aprobado en QA local')

# Sustituir solo las capturas de esta copia, conservando el documento aportado.
for col, api in enumerate((24, 35)):
    cell = doc.tables[5].rows[1].cells[col]
    cell.text = ''
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(EVIDENCE / f'api{api}/03-mascota-detalle.png'), width=Cm(6.0))
    for prop in para._p.xpath('.//wp:docPr'):
        prop.set('descr', f'Ficha ficticia guardada y consultada en Android API {api}')
cell_text(doc.tables[6].rows[4].cells[1], 'Se actualizaron los pendientes con pruebas reales automatizadas y revisión de capturas. Se conservan límites, incidencias y reflexiones personales por completar.')

for section in doc.sections:
    for para in section.footer.paragraphs:
        for run in para.runs:
            run.text = run.text.replace('Avance del 20 de octubre de 2026', 'Resultados del 22 de octubre de 2026')

def p(text):
    para = doc.add_paragraph(text)
    para.paragraph_format.space_before = Pt(5)
    return para
def heading(text):
    doc.add_page_break()
    doc.add_paragraph(text, 'Heading 1')
def data_table(headers, values):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:color'), 'D9D9D9')
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for cell, text in zip(table.rows[0].cells, headers): cell.text = text
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for values_row in values:
        for cell, text in zip(table.add_row().cells, values_row): cell.text = str(text)
    for i, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            if i == 0:
                shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'E7EDF1'); cell._tc.get_or_add_tcPr().append(shading)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(5)
                for run in para.runs:
                    run.font.size = Pt(9)
                    run.font.bold = i == 0
    return table

heading('9 Recuperación y observación de recursos')
p('En cada Android creamos una cuenta ficticia, cerramos sesión y volvimos a ingresar. Registramos una mascota Gato en comuna 2 y un reporte Encontrado en comuna 3. La consulta posterior mostró los datos guardados y las lecturas del servidor local confirmaron sus IDs y propietario.')
p('Después desactivamos solo la red del SDK de Firestore. El reporte nuevo quedó pendiente y el botón de envío se bloqueó. El verificador del servidor confirmó que aún no existía allí. Detuvimos el proceso con am force-stop, reabrimos y comprobamos el mismo ID. Al habilitar la red, el reporte se confirmó. El servidor terminó con una mascota y dos reportes por cuenta: el normal y el recuperado, sin duplicación.')
data_table(['Comprobación', 'Android 7 API 24', 'Android 15 API 35'], [
    ['Registro ingreso y cierre de sesión', 'Aprobado', 'Aprobado'],
    ['Alta consulta y datos de mascota', 'Aprobado', 'Aprobado'],
    ['Alta consulta y datos de reporte', 'Aprobado', 'Aprobado'],
    ['Pendiente ausente del servidor sin red', 'Comprobado', 'Comprobado'],
    ['Mismo ID tras reiniciar el proceso', 'Comprobado', 'Comprobado'],
    ['Servidor al terminar', '1 mascota y 2 reportes', '1 mascota y 2 reportes']
])
p('El intento de leer desde Android 15 un reporte de la cuenta usada en Android 7 fue rechazado con PERMISSION_DENIED. Las 33 pruebas de reglas pasaron antes de los recorridos. No se desplegaron reglas, índices ni datos remotos.')
mem_rows = []
for api, name in [(24, 'Android 7'), (35, 'Android 15')]:
    directory = median(runs[api]['directoryPssKiB']) / 1024
    map_pss = median(runs[api]['mapPssKiB']) / 1024
    mem_rows.append([name, f'{directory:.1f} MiB', f'{map_pss:.1f} MiB', f'+{map_pss-directory:.1f} MiB'])
data_table(['Entorno', 'Directorio mediana PSS', 'Mapa mediana PSS', 'Diferencia'], mem_rows)
p('PSS estima la memoria proporcional del proceso. Son tres muestras por pantalla, en orden directorio y después mapa, con la instrumentación activa y la caché existente. Los resultados pueden variar por carga, caché y sistema; no comparan toda la app antes y después. La prueba separada de bitmaps sí reproduce ambos muestreos con el decodificador real.')
p(f'SHA-256 del APK normal instalado en ambos Android: {normal_hash}')
p('El fallo de inicialización de QA se conserva como incidencia de preparación. Las ejecuciones finales taller22d terminaron aprobadas; su Logcat se delimitó por hora de inicio para no confundir mensajes históricos con el resultado actual.')

def image_page(title, pairs):
    heading(title)
    for caption, filename in pairs:
        para = p(caption)
        para.paragraph_format.keep_with_next = True
        table = doc.add_table(rows=1, cols=2)
        for col, api in enumerate((24,35)):
            cell = table.rows[0].cells[col]
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.add_run(f'Android {7 if api == 24 else 15} API {api}\n').font.size = Pt(9)
            para.add_run().add_picture(str(EVIDENCE / f'api{api}/{filename}.png'), width=Cm(4.7))
            for prop in para._p.xpath('.//wp:docPr'):
                prop.set('descr', f'{caption}, API {api}, datos ficticios')
        p('')

image_page('10 Sesión y consulta del reporte', [
    ('Sesión ciudadana después de salir y volver a ingresar', '01-sesion'),
    ('Reporte creado y consultado en la cuenta correspondiente', '05-reporte-detalle')
])
image_page('11 Pendiente recuperado y confirmación', [
    ('Pendiente conservado después de detener y reabrir el proceso', '10-pendiente-recuperado'),
    ('Confirmación del servidor local después de reconectar', '11-reconexion-confirmada')
])
for name in ('Title', 'Subtitle', 'Heading 1', 'Heading 2'):
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
# Mantener la tipografía del cuerpo aportado, también en los textos actualizados.
body_font = doc.paragraphs[5].runs[0].font.name or 'Times New Roman'
for para in list(doc.paragraphs) + [p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]:
    for run in para.runs:
        run.font.name = body_font
doc.core_properties.title = 'Compatibilidad Android de Mascotas Municipales'
doc.core_properties.subject = 'Resultados verificados en Android 7 y Android 15 con Firebase local'
doc.core_properties.author = 'Grupo Mascotas Municipales'
doc.core_properties.last_modified_by = 'Grupo Mascotas Municipales'
out = ROOT / 'documentos/Taller_Compatibilidad_Android_Resultados_2026-10-22.docx'
doc.save(out)
print(out)
